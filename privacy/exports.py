"""Local-only generators for watermarked educational privacy-export drafts."""

from __future__ import annotations

import hashlib
import json
import os
from dataclasses import dataclass
from pathlib import Path

from django.conf import settings
from django.utils import timezone


WATERMARK = "BẢN NHÁP HỌC TẬP - KHÔNG PHẢI HỒ SƠ NỘP CƠ QUAN NHÀ NƯỚC"


@dataclass(frozen=True)
class RenderedExport:
    file_name: str
    relative_name: str
    sha256: str


def _safe_filename(value: str, extension: str) -> str:
    name = Path(value).name.replace("\x00", "")
    stem = Path(name).stem or f"privacy-export-{timezone.now():%Y%m%d%H%M%S}"
    safe_stem = "".join(char if char.isalnum() or char in {"-", "_"} else "-" for char in stem).strip("-")
    return f"{safe_stem or 'privacy-export'}.{extension}"


def _export_directory() -> Path:
    root = Path(settings.EXPORT_ROOT).resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def _payload(artifact) -> dict[str, str]:
    case = artifact.case
    return {
        "watermark": WATERMARK,
        "generated_at": timezone.localtime().isoformat(),
        "organization": case.organization.name,
        "reference": case.reference,
        "request_type": case.get_request_type_display(),
        "case_status": case.get_status_display(),
        "subject_email": case.subject_email or "Đã ẩn danh hóa theo chính sách lưu giữ",
        "description": case.description or "Không có nội dung bổ sung.",
        "legal_notice": "Bản xuất minh họa cục bộ phục vụ seminar; không phải tư vấn pháp lý hoặc hồ sơ nộp cơ quan nhà nước.",
    }


def _render_docx(destination: Path, payload: dict[str, str]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = Document()
    watermark = document.add_paragraph()
    watermark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = watermark.add_run(WATERMARK)
    run.bold = True
    run.font.size = Pt(10)
    document.add_heading("Bản xuất yêu cầu quyền dữ liệu cá nhân", level=1)
    for label, key in (
        ("Tổ chức", "organization"),
        ("Mã hồ sơ", "reference"),
        ("Loại yêu cầu", "request_type"),
        ("Trạng thái", "case_status"),
        ("Chủ thể", "subject_email"),
        ("Thời điểm tạo", "generated_at"),
    ):
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(payload[key])
    document.add_heading("Nội dung yêu cầu", level=2)
    document.add_paragraph(payload["description"])
    document.add_paragraph(payload["legal_notice"])
    document.save(destination)


def _find_pdf_font() -> str | None:
    """Use a Unicode-capable font when present on Windows or Linux containers."""

    candidates = [
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "arial.ttf",
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            return str(candidate)
    return None


def _render_pdf(destination: Path, payload: dict[str, str]) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    font_name = "Helvetica"
    font_path = _find_pdf_font()
    if font_path:
        font_name = "PrivacyUnicode"
        pdfmetrics.registerFont(TTFont(font_name, font_path))

    page_width, page_height = A4
    pdf = canvas.Canvas(str(destination), pagesize=A4)
    pdf.setTitle("Educational Privacy Export Draft")
    pdf.setFont(font_name, 8)
    pdf.drawCentredString(page_width / 2, page_height - 35, WATERMARK)
    pdf.setFont(font_name, 16)
    pdf.drawString(48, page_height - 78, "Bản xuất yêu cầu quyền dữ liệu cá nhân")
    y = page_height - 112
    pdf.setFont(font_name, 10)
    for label, key in (
        ("Tổ chức", "organization"),
        ("Mã hồ sơ", "reference"),
        ("Loại yêu cầu", "request_type"),
        ("Trạng thái", "case_status"),
        ("Chủ thể", "subject_email"),
        ("Thời điểm tạo", "generated_at"),
    ):
        pdf.drawString(48, y, f"{label}: {payload[key]}")
        y -= 20
    y -= 8
    for heading, text in (("Nội dung yêu cầu", payload["description"]), ("Lưu ý", payload["legal_notice"])):
        pdf.setFont(font_name, 11)
        pdf.drawString(48, y, heading)
        y -= 18
        pdf.setFont(font_name, 9)
        # The generator intentionally keeps its own text compact and does not
        # attempt to reproduce a full legal notice or subject dataset.
        words = text.split()
        line = ""
        for word in words:
            proposed = f"{line} {word}".strip()
            if pdf.stringWidth(proposed, font_name, 9) > page_width - 96:
                pdf.drawString(48, y, line)
                y -= 14
                line = word
            else:
                line = proposed
        if line:
            pdf.drawString(48, y, line)
            y -= 22
    pdf.save()


def _render_json(destination: Path, payload: dict[str, str]) -> None:
    destination.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def render_educational_draft(artifact) -> RenderedExport:
    """Render a local draft and return only metadata needed for audit/storage.

    The caller is responsible for authorization, status transitions, and audit
    writes.  This function performs no network or regulatory filing action.
    """

    extension = artifact.format.lower()
    if extension not in {"pdf", "docx", "json"}:
        raise ValueError("Unsupported local export format.")
    file_name = _safe_filename(artifact.file_name, extension)
    destination = _export_directory() / file_name
    payload = _payload(artifact)
    if extension == "docx":
        _render_docx(destination, payload)
    elif extension == "pdf":
        _render_pdf(destination, payload)
    else:
        _render_json(destination, payload)
    digest = hashlib.sha256(destination.read_bytes()).hexdigest()
    return RenderedExport(file_name=file_name, relative_name=file_name, sha256=digest)
